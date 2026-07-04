from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "Documentation" / "Tasks" / "Task_1" / "Resources" / "CMSC_495_Phase_1_Template.docx"
OUTPUT = ROOT / "Documentation" / "Tasks" / "Task_1" / "Submissions" / "Phase_1_Source_Report_Final.docx"


def clear_document(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        # Keep section properties so page layout in the template remains valid.
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def add_cover_page(doc: Document) -> None:
    title = doc.add_paragraph("FocusFlow Productivity Application")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].bold = True
    title.runs[0].font.size = Pt(22)

    subtitle = doc.add_paragraph("CMSC 495 Phase I Source Code Report")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(16)

    doc.add_paragraph("")

    lines = [
        "Course: CMSC 495 Computer Science Capstone",
        "Team: Group Around the World",
        "Team Members: Danil Podolskiy, Jean Richard Fenelon, Stefan Nikolov, Alphonso Williams",
        f"Date: {date.today().isoformat()}",
    ]
    for line in lines:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    try:
        doc.add_heading(text, level=level)
    except KeyError:
        p = doc.add_paragraph(text)
        if p.runs:
            p.runs[0].bold = True
            p.runs[0].font.size = Pt(16 if level == 1 else 13)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(8)


def add_code_block(doc: Document, code: str) -> None:
    p = doc.add_paragraph(code)
    for run in p.runs:
        run.font.name = "Consolas"
        run.font.size = Pt(10)


def build_report(doc: Document) -> None:
    add_heading(doc, "Introduction", 1)
    add_body(
        doc,
        "FocusFlow is a Windows desktop productivity application developed for the CMSC 495 capstone. "
        "The product goal is to consolidate account management, task planning, calendar visibility, reminders, "
        "and performance tracking in one interface. During planning, the team identified a common productivity problem: "
        "users split their workflow across unrelated tools. That split creates context-switching overhead, inconsistent "
        "data entry, and weak visibility across deadlines and completed work. Phase I was therefore designed to establish "
        "a technically solid baseline that could support a complete working application in later phases."
    )
    add_body(
        doc,
        "The Phase I implementation now includes modular service-layer code, persistent storage bootstrapping, a professional "
        "desktop user interface, and structured test suites. The implementation is intentionally incremental: it validates "
        "the critical application backbone first, then layers interaction quality and persistence. This report documents the "
        "setup decisions, architecture choices, core features delivered, testing strategy, documentation practices, and how "
        "those decisions satisfy the grading rubric for project setup, core functionality, unit testing, and documentation."
    )

    add_heading(doc, "Objective", 1)
    add_body(
        doc,
        "The objective of this Phase I source code submission is to deliver an end-to-end baseline that proves the project is "
        "buildable, testable, and maintainable. The team targeted four outcomes mapped directly to the rubric categories:"
    )
    add_body(
        doc,
        "1) Project Setup: create a clean repository layout, dependency management, scripts, and source separation. "
        "2) Core Functionality: implement user authentication, task lifecycle controls, calendar task viewing, reminder CRUD, "
        "and dashboard statistics. 3) Unit Testing: create repeatable tests for major services and persistence adapters. "
        "4) Documentation: maintain readable code comments, complete run/build instructions and this formal report."
    )

    add_heading(doc, "Instructions", 1)

    add_heading(doc, "Review the Phase I", 2)
    add_body(
        doc,
        "Phase I requirements were reviewed from three sources: the project plan, project design, and assignment submission guidance. "
        "The plan and design established scope boundaries (desktop focus, core workflow priority, and security fundamentals), "
        "while the assignment guidance defined expectations for implementation depth and report quality. The team used those sources to derive "
        "a concrete execution list: establish development setup, build high-value business flows first, verify correctness with "
        "tests, and record rationale in reusable documentation."
    )
    add_body(
        doc,
        "The grading rubric emphasizes not only feature delivery but also engineering discipline. As a result, Phase I was "
        "organized into traceable workstreams. Project organization and automation satisfy setup quality. Service and UI work "
        "satisfy core functionality. Automated tests satisfy reliability and correctness. Structured runtime and submission documentation "
        "satisfy communication and maintainability. This review-first approach prevented ad hoc coding and improved alignment "
        "between code artifacts and grading criteria."
    )
    add_body(
        doc,
        "The team also used the rubric as a quality gate before each push. After every major implementation update, the code was "
        "evaluated against setup quality, functional scope, test coverage, and clarity of documentation. This process created a "
        "consistent acceptance standard that made peer review simpler and reduced subjective disagreement during merge decisions."
    )

    add_heading(doc, "Design the Software Architecture", 2)
    add_body(
        doc,
        "The implementation follows a layered architecture. The presentation layer is the PySide6 desktop interface. The business "
        "logic layer includes authentication, task management, and reminder management services. The data layer is repository-based, "
        "with in-memory adapters for rapid iteration and SQLAlchemy-backed adapters for persistent storage. This separation ensures "
        "that business rules remain testable independent of the GUI and that storage changes do not require UI rewrites."
    )
    add_body(
        doc,
        "Domain models include User, Task, and Reminder entities with typed status and priority enums. Task transitions are constrained "
        "through an explicit allowed-transition matrix to prevent invalid state changes. Reminders are tied to tasks and validated for "
        "ownership and scheduling correctness. Repository protocols define stable persistence contracts, and implementations for both "
        "in-memory and SQL persistence support development flexibility. This architecture directly supports modularity, minimizes "
        "duplication, and aligns with the rubric expectation for organized components and efficient implementation decisions."
    )
    add_body(
        doc,
        "A bootstrap layer initializes application services and selects the storage backend. By default, SQLite is used for local persistence "
        "to simplify setup for team members and graders. A database URL override supports MySQL-backed validation when needed. If a configured "
        "database is unavailable, the application can safely fall back to in-memory mode so demonstration flows remain available. This design "
        "balances reliability and practicality for an academic capstone timeline."
    )

    add_heading(doc, "Implement the Software", 2)
    add_body(
        doc,
        "Phase I implementation was completed in incremental milestones. First, the team created modular package structure, "
        "configured testing with pytest, and pinned dependencies in requirements.txt. Next, authentication and task modules were "
        "implemented with validation and custom exceptions. After core service behavior stabilized, persistent service bootstrapping was "
        "added with default SQLite storage and optional database URL override. Finally, the desktop interface was elevated from a static "
        "prototype to a professional tabbed experience."
    )
    add_body(
        doc,
        "The current interface includes account registration and login, task create/read/update/delete operations, status transitions, "
        "dashboard metrics, label summaries, calendar-based task viewing by selected date, and reminder create/update/acknowledge/delete "
        "workflows. UI layout was reorganized for readability with cards, tabs, consistent spacing, and grouped controls. The result is "
        "an application baseline that is functional for real usage and structured for future enhancements such as advanced analytics and "
        "expanded notification behavior."
    )
    add_body(
        doc,
        "Project cleanup and reorganization were also completed to improve long-term maintainability. Documentation was grouped into "
        "project references and task-specific resources/submissions. Naming conventions were normalized for clarity. Build and integration "
        "build and integration automation were consolidated. This cleanup addresses common capstone risks such as artifact sprawl, mixed naming, and "
        "difficulty onboarding teammates to file locations."
    )
    add_body(
        doc,
        "From a user-experience perspective, the desktop interface was reworked into a professional tabbed layout. The Tasks tab centralizes "
        "task lifecycle operations and editing. The Calendar tab displays tasks for selected dates to support planning. The Reminders tab "
        "supports create/update/acknowledge/delete workflows tied to existing tasks. Visual hierarchy, spacing, and control grouping were "
        "standardized to improve readability and reduce operator error when using the application under deadline conditions."
    )

    add_heading(doc, "Collaboration and Implementation", 2)
    add_body(
        doc,
        "Implementation was performed collaboratively with Git-based source control and small, reviewable commits. Major updates were grouped "
        "by objective: service foundations, persistence wiring, UI milestones, reminder/calendar expansion, and packaging support. This allowed "
        "the team to isolate regressions quickly and preserve a clear project history for checkpoint grading. Commit messages were kept explicit "
        "so reviewers could map each change to a project requirement or rubric category."
    )
    add_body(
        doc,
        "Code review focused on maintainability and behavior risk. Review comments emphasized input validation, ownership boundaries, transition "
        "rules, and test reliability. Where risks were detected, follow-up commits added or extended tests before feature expansion continued. "
        "This sequence (implement, test, review, refine) reduced rework and provided transparent evidence that quality controls were applied "
        "throughout implementation rather than deferred to the end of the phase."
    )
    add_body(
        doc,
        "Course resources and submission artifacts are now separated cleanly. Project reference documents are "
        "grouped by purpose (project references, provided resources, and submission outputs). This structure improves "
        "navigation during grading and helps future phases avoid accidental edits to instructor-provided templates or rubrics."
    )

    add_heading(doc, "Test the Software", 2)
    add_body(
        doc,
        "Testing is implemented using pytest and organized into unit and integration layers. Unit tests cover authentication rules, "
        "task workflow behavior, reminder service logic, and bootstrap behavior. Integration tests cover SQL repositories with environment-"
        "gated execution, ensuring database-backed operations are validated without forcing every local run to require a live DB instance. "
        "This two-layer strategy provides quick feedback during active development while preserving confidence in persistence behavior."
    )
    add_body(
        doc,
        "Test scenarios include successful and unsuccessful login, duplicate account prevention, task validation, status transition "
        "constraints, authorization boundaries, overdue processing, reminder scheduling constraints, reminder acknowledgment behavior, "
        "and repository CRUD round-trips. The team also executes smoke startup checks after major UI updates to ensure packaging and "
        "runtime pathways remain stable. This validation discipline reduces regressions and supports rubric criteria for comprehensive "
        "test coverage and clear testing framework usage."
    )
    add_body(
        doc,
        "Integration testing includes SQL repository round-trip validation and environment-gated execution for MySQL profiles. This ensures the "
        "same application logic can be validated under both local lightweight storage and relational database backends. The team also added a "
        "container-based test profile and runner scripts to keep database test setup reproducible. These controls improve confidence that the "
        "project can be demonstrated consistently across development environments."
    )
    add_body(
        doc,
        "Beyond automated tests, manual scenario checks were executed in the desktop UI for user workflows: account creation, login, task entry, "
        "task editing, status transitions, calendar date selection, reminder scheduling, reminder acknowledgment, and deletion flows. Manual "
        "checks focused on UX clarity and error handling messages, while automated tests focused on deterministic service behavior. Together, "
        "these methods provide both functional confidence and usability confidence for the current baseline."
    )

    add_heading(doc, "Project Setup Evidence", 1)
    add_body(
        doc,
        "Project setup quality is demonstrated by explicit tooling and structure decisions. The repository includes isolated source and test "
        "directories, deterministic dependency pins, startup entrypoints, bootstrap configuration controls, packaging scripts, and database test "
        "profiles. Build and run commands are documented in project run instructions. Cache and local artifact ignores are configured to keep commits clean and "
        "reviewable."
    )
    add_body(
        doc,
        "The project supports both development and demonstration modes. Developers can run directly from source with PYTHONPATH settings, while "
        "demonstrators can build a Windows executable through the packaging script. This dual-path setup ensures that capstone deliverables remain "
        "accessible to both technical and non-technical evaluators."
    )

    add_heading(doc, "Core Functionality Evidence", 1)
    add_body(
        doc,
        "Authentication flows now enforce required usernames, minimum password rules, duplicate account prevention, and secure password hashing. "
        "Task management supports create, load, edit, delete, status transitions, and overdue updates. Dashboard indicators summarize key workflow "
        "signals in one view. Calendar-based task viewing enables date-centric planning. Reminder workflows include create, update, acknowledge, "
        "and delete actions tied to task ownership validation."
    )
    add_body(
        doc,
        "These behaviors are implemented through modular services rather than embedded GUI logic. This service-first design allows functionality "
        "to remain testable and reusable if a future web or mobile interface is introduced. It also supports maintainable growth as additional "
        "features such as advanced reports, label management, and export capabilities are added in subsequent phases."
    )

    add_heading(doc, "Unit Testing Evidence", 1)
    add_body(
        doc,
        "Unit tests are organized by domain area: authentication, tasks, reminders, and bootstrap behavior. Integration tests validate repository "
        "persistence semantics under relational storage. The test suite includes happy-path and failure-path scenarios, which is essential for "
        "validation-heavy workflows like authentication and scheduling."
    )
    add_body(
        doc,
        "Examples of validated behaviors include prevention of short passwords, rejection of duplicate users, enforcement of valid task status "
        "transitions, authorization checks on cross-user access attempts, rejection of past reminder schedules, and confirmation of repository CRUD "
        "round-trips. This coverage demonstrates comprehensive attention to both correctness and misuse resistance."
    )

    add_heading(doc, "Documentation Evidence", 1)
    add_body(
        doc,
        "Documentation now exists at multiple levels. Inline docstrings communicate intent at module and class level. Run instructions describe storage and "
        "and packaging paths. Task artifacts are stored in structured folders for grader navigation. This report connects implementation choices to "
        "requirements and rubric expectations, creating an auditable record of engineering decisions."
    )
    add_body(
        doc,
        "The use of explicit section headings matching the provided template improves consistency with assignment expectations. References are "
        "included to support technical choices and course planning sources. Overall, documentation now supports maintenance, onboarding, and grading "
        "without requiring evaluators to reverse-engineer project intent from code alone."
    )
    add_body(
        doc,
        "To support reviewers who may only receive this document, implementation details are described directly in the report: architecture, workflow rules, validation behavior, test strategy, and representative executable examples. Evaluation does not depend on external code access."
    )

    add_heading(doc, "Risk Management and Technical Constraints", 1)
    add_body(
        doc,
        "Several technical constraints influenced implementation decisions in this phase. The first was environment variability: not every team "
        "member or grader will have an identical local stack. To mitigate this, the project defaults to SQLite while still supporting MySQL "
        "through a configurable connection string. This lowered setup risk without blocking persistence validation. A second constraint was UI "
        "complexity under limited timeline. The team prioritized professional readability and essential workflows over decorative complexity, which "
        "reduced delivery risk while still improving user experience quality."
    )
    add_body(
        doc,
        "A third risk involved feature coupling between GUI events and business rules. If UI code directly manages workflow logic, regression risk "
        "increases rapidly as features scale. The repository-service architecture mitigates this by centralizing validation and transition rules in "
        "testable services. A fourth risk was documentation drift where implementation changes are not reflected in assignment artifacts. The team "
        "addressed this by updating run instructions and submission documents in the same implementation cycle as feature updates."
    )
    add_body(
        doc,
        "Data integrity risk was also considered. Task and reminder ownership checks ensure that one user cannot manipulate another user's data. "
        "Reminder creation validates temporal constraints to block scheduling in the past. Task transition constraints prevent impossible workflow "
        "states. Together, these checks create a defensive baseline suitable for capstone demonstration while leaving room for tighter production-level "
        "controls such as role-based policy or audit logging in later phases."
    )

    add_heading(doc, "Verification Matrix", 1)
    add_body(
        doc,
        "To ensure broad quality coverage, the team mapped test intent to functional areas. Authentication tests validate registration rules, login "
        "behavior, and secure hashing verification. Task tests validate create/edit/delete operations, state transitions, ownership restrictions, and "
        "overdue handling. Reminder tests validate create/update/delete behavior, ownership enforcement, acknowledgment flow, and scheduling constraints. "
        "Bootstrap tests validate persistent startup behavior and fallback reliability."
    )
    add_body(
        doc,
        "Integration tests extend this matrix to persistence adapters. SQL-backed tests verify that user records, task records, and reminder records "
        "are correctly stored, retrieved, updated, and removed. These tests reduce risk that service behavior differs between in-memory and relational "
        "backends. Environment-gated execution was selected to avoid forcing heavy setup for every run while still allowing full verification in CI-like "
        "or instructor validation contexts."
    )
    add_body(
        doc,
        "Manual verification complements automated checks where visual confirmation is valuable. UI smoke checks validate that key views load, controls "
        "remain readable, and common workflows can be completed without hidden dependencies. This blended strategy reflects practical software engineering: "
        "automated tests provide deterministic confidence for logic, while manual checks validate usability and layout outcomes."
    )

    add_heading(doc, "Process Reflection and Lessons Learned", 1)
    add_body(
        doc,
        "One major lesson from this phase is that repository organization directly affects implementation speed. Before cleanup, task materials and "
        "reference files were mixed in flatter structures, which made artifact discovery slower. After reorganization into project references, task "
        "resources, and submission outputs, team coordination improved because expectations about file location became explicit."
    )
    add_body(
        doc,
        "A second lesson is that incremental commits with immediate testing significantly reduce integration pain. By validating after each major update, "
        "the team isolated bugs rapidly, including dependency mismatches and UI runtime issues. This approach preserved momentum and avoided large late-stage "
        "stabilization efforts. The process also produced a clearer commit history for grading, because each milestone maps to a tangible functional outcome."
    )
    add_body(
        doc,
        "A third lesson concerns balancing polish with scope. The team intentionally focused on professional readability and workflow clarity in the UI "
        "rather than attempting highly complex visual effects. This kept the interface reliable and maintainable while still delivering a polished user "
        "experience. In capstone contexts where grading emphasizes engineering quality and demonstrable outcomes, this tradeoff proved effective."
    )

    add_heading(doc, "Phase II Preparation", 1)
    add_body(
        doc,
        "With Phase I complete, Phase II can build on a stable baseline. Immediate next steps include deeper analytics (label/date trend summaries), "
        "expanded reminder behavior, and stronger packaging/distribution workflows. Because architecture boundaries are already established, these features "
        "can be added with lower refactor risk."
    )
    add_body(
        doc,
        "The team also plans to strengthen release readiness by adding explicit acceptance checklists for each milestone, improving integration test coverage "
        "for backend configurations, and continuing documentation synchronization with each merge. This will help ensure final submission quality is high "
        "not only in feature count but also in consistency, reliability, and maintainability."
    )

    add_heading(doc, "Rubric Alignment Summary", 1)
    add_body(
        doc,
        "Project Setup (25 points): The project now contains clear source/test/documentation separation, pinned dependencies, "
        "runtime bootstrap controls, version control history, and build scripts for Windows packaging. The environment setup is reproducible "
        "and documented in this report and companion run instructions."
    )
    add_body(
        doc,
        "Core Functionality (25 points): Delivered functionality includes registration/login, task CRUD, task transitions, dashboard metrics, "
        "calendar task visibility, reminder CRUD, and persistent storage integration. Code is modular and follows service-layer boundaries "
        "to reduce duplication."
    )
    add_body(
        doc,
        "Unit Testing (25 points): Multiple test modules cover business logic and persistence adapters. Tests are deterministic and include "
        "negative/edge-case validation. Integration tests are included for SQL-backed repositories and remain executable through configured "
        "test profiles."
    )
    add_body(
        doc,
        "Documentation (25 points): Codebase contains docstrings and structured exception messages; runtime and build details are documented; "
        "assignment documentation is reorganized into clear resources and submissions folders; this report provides traceability between objectives, "
        "implementation, testing, and grading criteria."
    )

    add_heading(doc, "Example Code", 1)
    add_body(doc, "The following sample illustrates a real Phase I workflow using the implemented services.")
    add_code_block(
        doc,
        "from datetime import date, datetime, timedelta, timezone\n"
        "from focusflow.bootstrap import create_service_bundle\n\n"
        "services = create_service_bundle()\n"
        "user = services.auth.register_user('student1', 'StrongPass123', 'student@example.com')\n"
        "task = services.tasks.create_task(\n"
        "    user_id=user.user_id,\n"
        "    title='Finalize capstone checkpoint',\n"
        "    due_date=date.today() + timedelta(days=2),\n"
        "    label='School'\n"
        ")\n"
        "from focusflow.models import TaskStatus\n"
        "services.tasks.transition_status(task_id=task.task_id, user_id=user.user_id, new_status=TaskStatus.IN_PROGRESS)\n"
        "services.reminders.create_reminder(\n"
        "    user_id=user.user_id,\n"
        "    task_id=task.task_id,\n"
        "    remind_at=datetime.now(timezone.utc) + timedelta(hours=4),\n"
        "    message='Review implementation and tests'\n"
        ")"
    )
    add_body(
        doc,
        "In application use, these operations are available through the desktop UI: users can authenticate, manage tasks, inspect schedules "
        "in the Calendar tab, and manage reminders in the Reminders tab. This code sample is included to demonstrate that the report is "
        "grounded in the implemented service contracts rather than abstract pseudocode."
    )

    add_heading(doc, "Conclusion", 1)
    add_body(
        doc,
        "Phase I successfully transforms FocusFlow from planning artifacts into a functioning software baseline. The application now includes "
        "a professional desktop user experience, modular business logic, persistent data pathways, robust validation, and meaningful test "
        "coverage. The repository cleanup completed during this cycle also improves maintainability and supports future team collaboration."
    )
    add_body(
        doc,
        "From a rubric perspective, the current submission demonstrates strong performance in all four grading categories: setup quality, core "
        "implementation depth, testing discipline, and documentation completeness. The next phase can focus on incremental enhancements (advanced "
        "statistics, richer notifications, and packaging polish) without reworking the underlying architecture."
    )
    add_body(
        doc,
        "Most importantly, this phase establishes a repeatable engineering process. The team can now plan future milestones against a stable baseline "
        "with known quality gates, clear ownership boundaries, and documented run/build instructions. That process maturity is a core capstone outcome "
        "and is critical to delivering a reliable final submission by the course deadline."
    )

    add_heading(doc, "References", 1)
    refs = [
        "GitHub Docs. (n.d.). About Projects. https://docs.github.com/issues/planning-and-tracking-with-projects/learning-about-projects/about-projects",
        "Group Around the World. (2026). Project Plan Group Around the World: FocusFlow [Course project plan]. University of Maryland Global Campus.",
        "OWASP Foundation. (n.d.). Password Storage Cheat Sheet. https://cheatsheetseries.owasp.org/cheatsheets/Password_Storage_Cheat_Sheet.html",
        "Oracle. (n.d.). MySQL 8.4 Reference Manual. https://dev.mysql.com/doc/refman/8.4/en/",
        "Qt Company. (n.d.). Qt for Python Documentation. https://doc.qt.io/qtforpython-6/",
        "SQLAlchemy. (n.d.). SQLAlchemy Documentation. https://docs.sqlalchemy.org",
    ]
    for ref in refs:
        add_body(doc, ref)


def main() -> None:
    doc = Document(str(TEMPLATE))
    clear_document(doc)
    add_cover_page(doc)
    add_heading(doc, "CMSC 495 Phase I Source Code", 1)
    build_report(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))


if __name__ == "__main__":
    main()
